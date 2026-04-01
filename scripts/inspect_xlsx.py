"""Temporary script to verify filenames in the Excel QA file match pdf_downloads."""

import os
import openpyxl

XLSX_PATH = "/Users/gregdunlop/projects2/statschat-ke/tests/accuracy/KNBS_Verified_QA_Examples_updated.xlsx"
PDF_DIR = "/Users/gregdunlop/projects2/statschat-ke/data/pdf_downloads"

pdf_files = set(os.listdir(PDF_DIR))

wb = openpyxl.load_workbook(XLSX_PATH)
ws = wb.active

headers = [cell.value for cell in ws[1]]
doc_col_idx = headers.index("relevant_doc_ids") + 1
ev_col_idx = headers.index("evidence_locations") + 1

print("=== Verification: do all doc_ids match files in pdf_downloads? ===\n")
all_ok = True
for row_num in range(2, ws.max_row + 1):
    qid = ws.cell(row=row_num, column=1).value
    doc_ids = ws.cell(row=row_num, column=doc_col_idx).value
    ev_locs = ws.cell(row=row_num, column=ev_col_idx).value

    if doc_ids:
        for doc in doc_ids.split(";"):
            doc = doc.strip()
            if doc and doc not in pdf_files:
                print(
                    f"  MISSING: Row {row_num} ({qid}) doc_id '{doc}' NOT in pdf_downloads/"
                )
                all_ok = False
            else:
                print(f"  OK: Row {row_num} ({qid}) doc_id '{doc}'")

    if ev_locs:
        for ev in ev_locs.split(";"):
            ev = ev.strip()
            # Extract filename (before page ref)
            fname = ev.split(" p.")[0].split(":p.")[0].strip()
            if fname and fname not in pdf_files:
                print(
                    f"  MISSING: Row {row_num} ({qid}) ev_loc file '{fname}' NOT in pdf_downloads/"
                )
                all_ok = False
            else:
                print(f"  OK: Row {row_num} ({qid}) ev_loc '{fname}'")

print()
if all_ok:
    print("ALL filenames match files in pdf_downloads/")
else:
    print("Some filenames do NOT match - see MISSING entries above")
